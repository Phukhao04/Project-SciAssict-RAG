"""
heading_heuristic.py
---------------------
Prototype fallback heading detector for docling_pipeline.py.

Problem it solves:
  Current pipeline reads heading level from the native Word style name
  (e.g. "Heading 1"). If the admin didn't click that style button, the
  paragraph is invisible as a heading -> chunking loses the
  "ปีที่ X / ภาคการศึกษาที่ Y" hierarchy that retrieval depends on.

What this does:
  1. First choice: still trust native style name if present (cheap, reliable).
  2. Fallback: if no native heading style, score the paragraph using
     formatting signals that headings *naturally* have even when nobody
     applied a style - larger font size than body text, bold, short text,
     no trailing punctuation.
  3. Every fallback classification carries a confidence score. Low-confidence
     hits are exactly what should surface in the admin chunk-review UI
     discussed for the auto-tag + editable-chunk feature, instead of
     silently trusting or silently discarding them.

This is a standalone prototype - not wired into docling_pipeline.py yet,
so you can see the logic and output shape before deciding where to splice
it into the real HierarchicalChunker flow.
"""

from __future__ import annotations
from dataclasses import dataclass, field
from collections import Counter
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
import re


# --------------------------------------------------------------------------
# 1. Establish the document's "body text" baseline font size.
#    Headings only make sense relative to this - a 14pt heading in a
#    document whose body is 8pt is a very different signal than 14pt
#    in a document whose body is 12pt.
# --------------------------------------------------------------------------

DEFAULT_BODY_SIZE_PT = 11.0  # fallback if we truly can't detect anything


def resolve_run_size_pt(run, paragraph) -> float:
    """A run's real rendered size is often NOT set on the run itself -
    most body text just inherits it from the paragraph style (or that
    style's base_style chain). Only walking this chain gives an accurate
    picture; reading run.font.size alone silently mistakes "unstyled body
    text" for "no signal" and skews the whole baseline."""
    if run.font.size is not None:
        return round(run.font.size.pt, 1)
    style = paragraph.style
    while style is not None:
        if style.font.size is not None:
            return round(style.font.size.pt, 1)
        style = style.base_style
    return DEFAULT_BODY_SIZE_PT


def get_body_font_size(doc: Document) -> float:
    sizes = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        # Paragraphs that already carry a native Heading style must NOT
        # count toward the body baseline - otherwise a document that mixes
        # correctly-styled headings with forgotten-style ones corrupts the
        # very thing we're trying to measure.
        if _native_style_level(para) is not None:
            continue
        runs_with_text = [r for r in para.runs if r.text.strip()]
        for run in runs_with_text:
            sizes.append(resolve_run_size_pt(run, para))
    if not sizes:
        return DEFAULT_BODY_SIZE_PT
    # Body text is whatever size appears most often - headings are the
    # minority by definition.
    most_common_size, _ = Counter(sizes).most_common(1)[0]
    return most_common_size


# --------------------------------------------------------------------------
# 2. Per-paragraph classification
# --------------------------------------------------------------------------

@dataclass
class HeadingResult:
    text: str
    is_heading: bool
    level: int | None          # 1 = top level, 2 = sub, 3 = sub-sub
    confidence: float          # 0.0 - 1.0
    source: str                # "native_style" | "heuristic" | "body"
    reasons: list[str] = field(default_factory=list)


NATIVE_HEADING_RE = re.compile(r"^Heading\s*(\d)$", re.IGNORECASE)

# Headings are almost never long sentences - character count is used
# instead of word count because Thai script has no spaces between words,
# so a naive word-split makes a whole Thai sentence look like "2 words".
MAX_HEADING_CHARS = 50
# Trailing-punctuation is a weak signal for Thai specifically (Thai
# sentences rarely end in "." the way English ones do), so it gets a
# much smaller weight than the size/bold signals below.
TRAILING_PUNCT = (".", "। ", "ๆ.", "!")


def _native_style_level(para: Paragraph) -> int | None:
    m = NATIVE_HEADING_RE.match(para.style.name or "")
    return int(m.group(1)) if m else None


def _paragraph_max_size(para: Paragraph, body_pt: float) -> float:
    runs_with_text = [r for r in para.runs if r.text.strip()]
    if not runs_with_text:
        return body_pt
    return max(resolve_run_size_pt(r, para) for r in runs_with_text)


def _paragraph_is_bold(para: Paragraph) -> bool:
    runs_with_text = [r for r in para.runs if r.text.strip()]
    if not runs_with_text:
        return False
    bold_runs = sum(1 for r in runs_with_text if r.bold)
    return bold_runs / len(runs_with_text) >= 0.8  # mostly-bold paragraph


def classify_paragraph(para: Paragraph, body_pt: float) -> HeadingResult:
    text = para.text.strip()

    # --- Path A: native style wins outright, high confidence -------------
    native_level = _native_style_level(para)
    if native_level is not None:
        return HeadingResult(
            text=text, is_heading=True, level=native_level,
            confidence=1.0, source="native_style",
            reasons=[f"Word style '{para.style.name}'"],
        )

    if not text:
        return HeadingResult(text="", is_heading=False, level=None,
                              confidence=1.0, source="body", reasons=["empty"])

    # --- Path B: heuristic fallback ---------------------------------------
    size = _paragraph_max_size(para, body_pt)
    bold = _paragraph_is_bold(para)
    short = len(text) <= MAX_HEADING_CHARS
    no_trailing_punct = not text.rstrip().endswith(TRAILING_PUNCT)
    size_ratio = size / body_pt if body_pt else 1.0

    signals = []
    score = 0.0

    if size_ratio >= 1.5:
        signals.append(f"font {size:.0f}pt is {size_ratio:.1f}x body ({body_pt:.0f}pt)")
        score += 0.50
    elif size_ratio >= 1.15:
        signals.append(f"font {size:.0f}pt is {size_ratio:.1f}x body ({body_pt:.0f}pt)")
        score += 0.30
    elif size_ratio > 1.0:
        signals.append(f"font {size:.0f}pt is slightly above body ({body_pt:.0f}pt)")
        score += 0.10

    if bold:
        signals.append("mostly bold run(s)")
        score += 0.25

    if short:
        signals.append(f"short ({len(text)} chars)")
        score += 0.15

    if no_trailing_punct:
        signals.append("no trailing punctuation")
        score += 0.05

    # Hard gate: text rendered at the SAME size as body can basically never
    # be a structural heading no matter how bold - that combination is
    # overwhelmingly just emphasis (e.g. a bolded warning sentence), not
    # section structure. Only strictly-larger-than-body text is eligible.
    is_heading = score >= 0.5 and size_ratio > 1.0
    if not is_heading:
        return HeadingResult(text=text, is_heading=False, level=None,
                              confidence=1.0 - score, source="body",
                              reasons=["below heading threshold"])

    # crude level guess from how far above body size it sits
    if size_ratio >= 1.5 and bold:
        level = 1
    elif size_ratio >= 1.15:
        level = 2
    else:
        level = 3

    return HeadingResult(text=text, is_heading=True, level=level,
                          confidence=round(min(score, 0.95), 2),  # cap: never as sure as native
                          source="heuristic", reasons=signals)


# --------------------------------------------------------------------------
# 3. Walk the document IN ORDER (paragraphs + tables interleaved) -
#    reusing the lesson already learned: doc.paragraphs / doc.tables as
#    separate lists destroys the heading -> table relationship.
# --------------------------------------------------------------------------

def iter_block_items(doc: Document):
    """Yield paragraphs and tables in the order they actually appear."""
    from docx.oxml.ns import qn
    parent_elm = doc.element.body
    for child in parent_elm.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, doc)
        elif child.tag == qn('w:tbl'):
            yield Table(child, doc)


def detect_structure(docx_path: str) -> list[dict]:
    doc = Document(docx_path)
    body_pt = get_body_font_size(doc)

    results = []
    for block in iter_block_items(doc):
        if isinstance(block, Table):
            results.append({
                "type": "table",
                "preview": " | ".join(c.text.strip() for c in block.rows[0].cells) if block.rows else "",
                "level": None, "confidence": None, "source": "table",
            })
            continue

        r = classify_paragraph(block, body_pt)
        if not r.text:
            continue
        results.append({
            "type": "heading" if r.is_heading else "body",
            "preview": r.text[:60],
            "level": r.level,
            "confidence": r.confidence,
            "source": r.source,
            "reasons": r.reasons,
        })
    return results


if __name__ == "__main__":
    import sys
    path = sys.argv[1] if len(sys.argv) > 1 else "sample_no_native_headings.docx"
    for row in detect_structure(path):
        tag = f"[{row['type'].upper():6}]"
        lvl = f"L{row['level']}" if row.get("level") else "  "
        conf = f"{row['confidence']:.2f}" if row.get("confidence") is not None else "  - "
        print(f"{tag} {lvl:3} conf={conf} src={row['source']:12} | {row['preview']}")